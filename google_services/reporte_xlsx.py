"""
El reporte de vínculos en Excel, para el botón de descarga del panel.

Tres hojas: las cifras, el listado completo y las filas sospechosas. Es el mismo
contenido que el CSV más el contexto que un CSV no puede llevar (varias tablas,
formato, y la frase que dice qué recorte se está mirando).

POR QUÉ EN MODO `write_only`
    openpyxl normal mantiene el libro entero en memoria: medido sobre las 23 620
    filas reales, 61,9 MB de pico frente a 1,1 MB en `write_only`. Con 3 workers x
    8 hilos, tres descargas a la vez se comerían ~190 MB de un worker. El modo de
    escritura secuencial vuelca las filas a un temporal según llegan.

    El precio es que NO se pueden combinar celdas ni volver atrás: las hojas se
    escriben de arriba abajo y de una sola pasada. Por eso el resumen son filas
    etiquetadas y no tarjetas.

El libro se arma en memoria y se devuelve entero (~0,7 MB): un .xlsx es un ZIP y
no se puede empezar a enviar antes de cerrarlo, así que esto no puede ir en
streaming como el CSV.
"""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Paleta: la misma que los informes de herramientas (informe_gerencial_google.py).
# Se repite aquí a propósito en vez de importarla: `google_services` es un servicio
# y no puede depender de un script suelto de la raíz.
TINTA = "0B0B0B"
TINTA_2 = "52514E"
BUENO = "0CA30C"
AVISO = "FAB219"
CRITICO = "D03B3B"
AZUL = "2A78D6"
GRIS = "898781"
FONDO_CABECERA = "F0EFEC"

ORIGEN_TEXTO = {
    "creacion": ("Creadas por la API", BUENO),
    "vinculacion": ("Ya existían; la API les puso la cédula", AZUL),
    "backfill": ("Sembradas por el backfill", AZUL),
    "sincronizacion": ("Ya existían y ya tenían cédula", GRIS),
    "manual": ("Vinculadas a mano", AVISO),
}

ANOMALIA_TEXTO = {
    "correo_vacio": ("Correo vacío", CRITICO,
                     "La fila no tiene dirección; no se puede notificar a esa persona."),
    "dominio_ajeno": ("Correo de otro dominio", AVISO,
                      "La dirección no es del dominio ni de un subdominio suyo."),
    "sin_google_id": ("Sin identificador de Google", CRITICO,
                      "Sin google_id no se puede volver a encontrar la cuenta."),
    "cedula_de_relleno": ("Cédula de relleno", CRITICO,
                          "Vacía o un dígito repetido: la fila no identifica a nadie."),
    "correo_repetido": ("Correo repetido", AVISO,
                        "La misma dirección aparece en más de una fila."),
}

# (clave en la fila, encabezado, ancho). El orden manda: es el de las celdas.
COLUMNAS = (
    ("identificacion", "Cédula / documento", 20),
    ("email", "Correo institucional", 34),
    ("ou", "Unidad organizativa", 30),
    ("principal", "Principal", 11),
    ("consumidor", "Sistema que lo registró", 22),
    ("origen", "Origen", 16),
    ("creado_en", "Registrado", 20),
    ("actualizado_en", "Actualizado", 20),
    ("google_id", "ID de Google", 24),
)


def _celda(ws, valor, *, negrita=False, color=TINTA, tam=10, fondo=None, ajuste=False):
    c = WriteOnlyCell(ws, value=valor)
    c.font = Font(size=tam, bold=negrita, color=color)
    if fondo:
        c.fill = PatternFill("solid", fgColor=fondo)
    if ajuste:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def _anchos(ws, anchos):
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a


def _hoja_resumen(wb, resumen: dict, descripcion: str):
    ws = wb.create_sheet("Resumen")
    _anchos(ws, (38, 26, 12, 12, 14, 14))

    ws.append([_celda(ws, "Correos institucionales registrados", negrita=True, tam=15)])
    ws.append([_celda(ws, descripcion, color=TINTA_2)])
    ws.append([_celda(ws, "Sale de google_vinculos, el índice LOCAL de las cuentas de "
                          "Google. Que una cuenta siga viva en el directorio lo "
                          "responde /personas/{cedula}/confirmar, no este reporte.",
                      color=TINTA_2)])
    ws.append([])

    # Creadas y adoptadas, separadas: una cuenta que ya existía y solo recibió la
    # cédula no es una cuenta creada.
    por_origen = {}
    for f in resumen.get("por_origen", []):
        por_origen[f["origen"]] = por_origen.get(f["origen"], 0) + f["cuentas"]

    ws.append([_celda(ws, "Cifras de la tabla completa", negrita=True, tam=12)])
    for etiqueta, valor, color in (
        ("Cuentas registradas", resumen.get("vinculos", 0), TINTA),
        ("Personas distintas", resumen.get("personas", 0), TINTA),
        ("Creadas por la API", por_origen.get("creacion", 0), BUENO),
        ("Ya existían; la API les puso la cédula", por_origen.get("vinculacion", 0), AZUL),
        ("Tipos de anomalía", len(resumen.get("anomalias", [])),
         CRITICO if resumen.get("anomalias") else GRIS),
    ):
        ws.append([_celda(ws, etiqueta, color=TINTA_2),
                   _celda(ws, valor, negrita=True, color=color, tam=12)])
    ws.append([])

    ws.append([_celda(ws, "Por origen y sistema", negrita=True, tam=12)])
    ws.append([_celda(ws, t, negrita=True, fondo=FONDO_CABECERA)
               for t in ("Origen", "Sistema", "Cuentas", "Personas", "Primera", "Última")])
    orden = list(ORIGEN_TEXTO)
    filas = sorted(resumen.get("por_origen", []),
                   key=lambda f: (orden.index(f["origen"]) if f["origen"] in orden
                                  else len(orden), -f["cuentas"]))
    for f in filas:
        etiqueta, color = ORIGEN_TEXTO.get(f["origen"], (f["origen"], TINTA_2))
        ws.append([_celda(ws, etiqueta, negrita=True, color=color),
                   _celda(ws, f["consumidor"], color=TINTA_2),
                   _celda(ws, f["cuentas"]), _celda(ws, f["personas"]),
                   _celda(ws, f["primera"], color=TINTA_2),
                   _celda(ws, f["ultima"], color=TINTA_2)])
    ws.append([])

    por_dia = resumen.get("por_dia", [])
    ws.append([_celda(ws, "Altas por día (ventana del resumen)", negrita=True, tam=12)])
    if not por_dia:
        ws.append([_celda(ws, "Sin altas en el periodo.", color=TINTA_2)])
        return
    ws.append([_celda(ws, t, negrita=True, fondo=FONDO_CABECERA)
               for t in ("Día", "Origen", "Sistema", "Altas")])
    for f in por_dia:
        etiqueta, color = ORIGEN_TEXTO.get(f["origen"], (f["origen"], TINTA_2))
        ws.append([_celda(ws, f["dia"]), _celda(ws, etiqueta, color=color),
                   _celda(ws, f["consumidor"], color=TINTA_2), _celda(ws, f["altas"])])


def _hoja_correos(wb, filas) -> int:
    """Escribe el listado según llega del generador. Devuelve cuántas filas puso."""
    ws = wb.create_sheet("Correos")
    _anchos(ws, [c[2] for c in COLUMNAS])
    ws.freeze_panes = "A2"
    ws.append([_celda(ws, c[1], negrita=True, fondo=FONDO_CABECERA) for c in COLUMNAS])

    n = 0
    for fila in filas:
        celdas = []
        for clave, _, _ in COLUMNAS:
            valor = fila.get(clave)
            if clave == "principal":
                valor = "Sí" if valor else "No"
            celdas.append(_celda(ws, "" if valor is None else valor))
        ws.append(celdas)
        n += 1

    # El autofiltro se declara al final porque necesita el nº de filas, pero no
    # obliga a recorrerlas otra vez: es solo una referencia.
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{max(n + 1, 1)}"
    return n


def _hoja_anomalias(wb, resumen: dict):
    ws = wb.create_sheet("Anomalías")
    _anchos(ws, (28, 10, 54, 46))

    ws.append([_celda(ws, "Filas sospechosas", negrita=True, tam=13)])
    ws.append([_celda(ws, "Defectos de FORMA detectables desde la base. NO dicen si la "
                          "cuenta sigue existiendo en Google.", color=TINTA_2)])
    ws.append([])

    anomalias = resumen.get("anomalias", [])
    if not anomalias:
        ws.append([_celda(ws, "Ninguna. Todas las filas tienen buena forma.",
                          negrita=True, color=BUENO, tam=11)])
        return

    ws.append([_celda(ws, t, negrita=True, fondo=FONDO_CABECERA)
               for t in ("Problema", "Filas", "Qué significa", "Ejemplos")])
    for a in anomalias:
        etiqueta, color, explicacion = ANOMALIA_TEXTO.get(
            a["problema"], (a["problema"], AVISO, ""))
        ws.append([_celda(ws, etiqueta, negrita=True, color=color),
                   _celda(ws, a["filas"], negrita=True, color=color),
                   _celda(ws, explicacion, color=TINTA_2, ajuste=True),
                   _celda(ws, ", ".join(a.get("ejemplos") or []), color=TINTA_2,
                          tam=9, ajuste=True)])


def generar(filas, resumen: dict, descripcion: str) -> bytes:
    """
    Arma el libro y lo devuelve en bytes.

    `filas` es el generador de `vinculos.exportar()`: se consume una sola vez, en
    orden, sin materializarlo. `resumen` es lo que devuelve `vinculos.resumen()`.
    `descripcion` es la frase que explica qué filtro se aplicó.
    """
    wb = Workbook(write_only=True)
    _hoja_resumen(wb, resumen, descripcion)
    _hoja_correos(wb, filas)
    _hoja_anomalias(wb, resumen)

    memoria = io.BytesIO()
    wb.save(memoria)
    return memoria.getvalue()


def nombre_archivo() -> str:
    return f"correos_{date.today().isoformat()}.xlsx"
